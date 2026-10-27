"""Extract public PDF/XLSX fee tables without losing original binary evidence."""
import io, re, shutil, subprocess
from html import escape


def image_html(raw):
    """Extract bounded Chinese/English text from an image when Tesseract exists."""
    if not shutil.which('tesseract'): return None
    try:
        result=subprocess.run(['tesseract','stdin','stdout','-l','chi_sim+eng'],
            input=raw,capture_output=True,timeout=30)
    except (OSError,subprocess.TimeoutExpired):
        return None
    text=result.stdout.decode('utf-8',errors='replace').strip()
    return '<main><pre>'+escape(text)+'</pre></main>' if text else None


def document_html(raw,content_type,url):
    if raw.startswith(b'%PDF') or 'application/pdf' in content_type:
        import pdfplumber
        tables=[];texts=[];previous_header=None
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            for page in pdf.pages[:35]:
                texts.append(page.extract_text(layout=False) or '')
                for table in page.extract_tables():
                    # PDF tables often continue on the next page without repeating the header.
                    cleaned=[[re.sub(r'\s+','',str(v or '')) for v in row] for row in table]
                    own=next((r for r in cleaned[:4] if any('手续费' in v for v in r) and any(v in ('品种','合约','代码','品种代码','合约名称') for v in r)),None)
                    if own:previous_header=own
                    elif previous_header and cleaned and len(cleaned[0])==len(previous_header):
                        code_cols=[i for i,v in enumerate(previous_header) if v in ('代码','品种代码','品种')]
                        if any(any(re.fullmatch(r'[a-zA-Z]{1,3}(?:[-0-9a-zA-Z/]+)?',r[i]) for i in code_cols) for r in cleaned[:3]):
                            table=[previous_header]+table
                    tables.append(table)
        html='<main><pre>'+escape('\n'.join(texts))+'</pre>'
        for table in tables:
            html+='<table>'+''.join('<tr>'+''.join('<td>'+escape(str(v or ''))+'</td>' for v in row)+'</tr>' for row in table)+'</table>'
        return html+'</main>',('PDF' if any(t.strip() for t in texts) or tables else 'PDF扫描件'),texts
    if raw.startswith(b'PK') and ('.xlsx' in url.lower() or 'spreadsheet' in content_type):
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        html='';texts=[]
        for sheet in book:
            html+='<table>'
            for row in sheet.iter_rows(max_row=min(sheet.max_row,2000),values_only=True):
                cells=[str(v) if v is not None else '' for v in row]
                texts.append(' | '.join(cells))
                html+='<tr>'+''.join('<td>'+escape(v)+'</td>' for v in cells)+'</tr>'
            html+='</table>'
        book.close();return html,'Excel',texts
    return None,None,[]
